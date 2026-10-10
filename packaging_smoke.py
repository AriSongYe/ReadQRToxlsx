"""Frozen-app smoke test using generated data only; never opens user files."""
from datetime import datetime
import json
from pathlib import Path
import tempfile
import traceback
from zipfile import ZipFile


def run(output_path):
    report = {"ok": False, "checks": []}
    try:
        import tkinter as tk
        import openpyxl
        from PIL import Image
        import zxingcpp
        from QRParsing import decode_image
        from UILayout import MyApp
        from readxlsx import ExcelManager

        root = tk.Tk()
        root.withdraw()
        app = MyApp(root)
        root.update()
        app._destroy()
        report["checks"].append("Tk UI initialization")
        with tempfile.TemporaryDirectory() as directory:
            directory = Path(directory)
            image_path = directory / "qr.png"
            barcode = zxingcpp.create_barcode("1-7", zxingcpp.BarcodeFormat.QRCode)
            Image.fromarray(barcode.to_image(scale=5)).save(image_path)
            assert decode_image(image_path) == ["1-7"]
            report["checks"].append("Pillow and zxing-cpp native QR decoding")
            workbook_path = directory / "smoke.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "OFF선교헌금"
            sheet.append(["일반QR", "신학생QR", "선교사QR", "개척교회QR", "북한QR", "이름"])
            sheet.cell(1, 7, "2026년 작정금액")
            sheet.cell(1, 12, datetime(2026, 10, 4))
            for category, label in enumerate(["일반", "신학생", "선교사", "개척", "북한"], 1):
                sheet.cell(2, 6 + category, label)
                sheet.cell(2, 11 + category, label)
                sheet.cell(2401, category, f"{category}-7")
                sheet.cell(2401, 6 + category, category * 1000)
            sheet.cell(2401, 6, "테스트")
            workbook.save(workbook_path)
            workbook.close()
            manager = ExcelManager(workbook_path)
            day = datetime(2026, 10, 4).date()
            for category in range(1, 6):
                item = manager.lookup(f"{category}-7", 2026, day)
                assert item.row == 2401 and item.pledged == category * 1000
            item = manager.lookup("1-7", 2026, day)
            manager.record(item, 3000)
            backup = manager.save()
            with ZipFile(backup) as before, ZipFile(workbook_path) as after:
                changed = {name for name in before.namelist() if before.read(name) != after.read(name)}
                assert changed == {"xl/workbook.xml", "xl/worksheets/sheet1.xml"}
            manager.close()
            reopened = ExcelManager(workbook_path)
            assert reopened.lookup("1-7", 2026, day).existing == 3000
            reopened.close()
            report["checks"].append("XLSX lookup, backup, save and reopen")
        report["ok"] = True
    except Exception:
        report["error"] = traceback.format_exc()
    Path(output_path).write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return 0 if report["ok"] else 1
