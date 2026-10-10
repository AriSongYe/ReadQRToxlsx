"""Financial write guards and regression checks; supplied files stay untouched."""
from datetime import date, datetime
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill

from QRParsing import parse_qr
from readxlsx import ExcelManager, parse_amount

ROOT = Path(__file__).resolve().parents[1]
DAY = date(2026, 10, 4)


def fixture(path):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "OFF선교헌금 "
    sheet.append(["일반QR", "신학생QR", "선교사QR", "개척교회QR", "북한QR", "id", "이름"])
    for start, header in [(8, "2025년 작정금액"), (13, "2026년 작정금액"), (18, datetime(2026, 10, 4)), (23, datetime(2025, 10, 5))]:
        sheet.cell(1, start, header)
        for offset, category in enumerate(["일반", "신학생", "선교사", "개척", "북한"]):
            sheet.cell(2, start + offset, category)
    for row, identifier in [(3, 1), (2401, 7007)]:
        for category in range(1, 6):
            sheet.cell(row, category, f"{category}-{identifier}")
            sheet.cell(row, 12 + category, 1000 * category)
            sheet.cell(row, 7 + category, 500 * category)
        sheet.cell(row, 7, f"교인{identifier}")
    sheet["R2401"].fill = PatternFill("solid", fgColor="00FF00")
    sheet["R2401"].number_format = '#,##0'
    sheet["M3"].comment = Comment("보존할 주석", "User")
    sheet["AB3"] = "=SUM(M3:Q3)"
    workbook.save(path)
    workbook.close()


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.path = Path(self.temp.name) / "test.xlsx"
        fixture(self.path)
        self.manager = ExcelManager(self.path)

    def tearDown(self):
        self.manager.close()
        self.temp.cleanup()

    def test_all_categories_beyond_old_row_limit(self):
        for category in range(1, 6):
            found = self.manager.lookup(f"{category}-7007\r\n", 2026, DAY)
            self.assertEqual(found.row, 2401)
            self.assertEqual(found.pledged, 1000 * category)
            self.assertEqual(found.coordinate, f"{'RSTUV'[category-1]}2401")

    def test_year_and_date_validation(self):
        self.assertEqual(self.manager.years, [2025, 2026])
        self.assertEqual(self.manager.lookup("1-1", 2025, date(2025, 10, 5)).pledged, 500)
        for year, day in [(2024, DAY), (2025, DAY), (2026, date(2026, 10, 11))]:
            with self.assertRaises(ValueError):
                self.manager.lookup("1-1", year, day)

    def test_invalid_and_missing_qr(self):
        for text in ["", "6-1", "1", "1-", "https://example.org", "1-x", "1-1-2"]:
            with self.assertRaises(ValueError):
                parse_qr(text)
        with self.assertRaises(ValueError):
            self.manager.lookup("1-999999", 2026, DAY)

    def test_duplicate_qr_is_blocked(self):
        self.manager._indexes[self.manager.sheet_name]["1-1"].append(2401)
        with self.assertRaisesRegex(ValueError, "중복"):
            self.manager.lookup("1-1", 2026, DAY)

    def test_amount_and_overwrite_guards(self):
        self.assertEqual(parse_amount("3,000"), 3000)
        self.assertEqual(parse_amount("0"), 0)
        for text in ["", "-1", "1.5", "abc", "1000000000000000"]:
            with self.assertRaises(ValueError):
                parse_amount(text)
        donation = self.manager.lookup("1-7007", 2026, DAY)
        self.manager.record(donation, 3000)
        with self.assertRaises(ValueError):
            self.manager.record(donation, 4000)
        self.manager.record(donation, 0, overwrite=True)
        self.assertEqual(self.manager.lookup("1-7007", 2026, DAY).existing, 0)

    def test_save_backup_style_formula_and_archive_preservation(self):
        original = self.path.read_bytes()
        self.manager.record(self.manager.lookup("1-7007", 2026, DAY), 3000)
        # S3 does not exist in original XML: exercise ordered cell insertion.
        self.manager.record(self.manager.lookup("2-1", 2026, DAY), 2000)
        backup = self.manager.save()
        self.assertEqual(backup.read_bytes(), original)
        self.assertFalse(self.manager.dirty)
        self.assertIsNone(self.manager.save())
        with ZipFile(backup) as before, ZipFile(self.path) as after:
            self.assertEqual(before.namelist(), after.namelist())
            changed = [name for name in before.namelist() if before.read(name) != after.read(name)]
            self.assertEqual(set(changed), {"xl/workbook.xml", "xl/worksheets/sheet1.xml"})
        workbook = openpyxl.load_workbook(self.path)
        sheet = workbook["OFF선교헌금 "]
        self.assertEqual(sheet["R2401"].value, 3000)
        self.assertEqual(sheet["S3"].value, 2000)
        self.assertEqual(sheet["R2401"].fill.fgColor.rgb, "0000FF00")
        self.assertEqual(sheet["R2401"].number_format, '#,##0')
        self.assertEqual(sheet["AB3"].value, "=SUM(M3:Q3)")
        self.assertEqual(sheet["M3"].comment.text, "보존할 주석")
        self.assertTrue(workbook.calculation.fullCalcOnLoad)
        workbook.close()
        # Saved values remain protected after the next save, not only in staging.
        with self.assertRaises(ValueError):
            self.manager.record(self.manager.lookup("1-7007", 2026, DAY), 5000)

    def test_external_change_and_failed_replace_preserve_original(self):
        self.manager.record(self.manager.lookup("1-7007", 2026, DAY), 3000)
        original = self.path.read_bytes()
        with patch("readxlsx.os.replace", side_effect=PermissionError("Excel open")):
            with self.assertRaises(PermissionError):
                self.manager.save()
        self.assertEqual(self.path.read_bytes(), original)
        self.assertTrue(self.manager.dirty)
        self.assertFalse(list(self.path.parent.glob(".qr-save-*")))
        self.path.write_bytes(original + b"external change")
        with self.assertRaisesRegex(ValueError, "外部|외부"):
            self.manager.save()
        self.assertEqual(self.path.read_bytes(), original + b"external change")

    def test_formula_destination_is_blocked(self):
        self.manager.close()
        workbook = openpyxl.load_workbook(self.path)
        workbook["OFF선교헌금 "]["R2401"] = "=1000+2000"
        workbook.save(self.path)
        workbook.close()
        self.manager = ExcelManager(self.path)
        with self.assertRaisesRegex(ValueError, "수식"):
            self.manager.record(self.manager.lookup("1-7007", 2026, DAY), 3000)




if __name__ == "__main__":
    unittest.main()
