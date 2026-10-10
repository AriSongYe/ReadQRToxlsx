"""Indexed lookups and surgical XLSX writes preserving untouched ZIP members."""
from dataclasses import dataclass
from datetime import date, datetime
from hashlib import sha256
from io import BytesIO
from pathlib import Path
import os
import posixpath
import re
import shutil
import tempfile
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.utils.cell import coordinate_from_string
from QRParsing import CATEGORIES, parse_qr

MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def parse_amount(value):
    text = str(value).strip().replace(",", "")
    if not re.fullmatch(r"[0-9]+", text):
        raise ValueError("금액은 0 이상의 정수로 입력해 주세요. 예: 3,000")
    amount = int(text)
    if amount > 999_999_999_999_999:
        raise ValueError("금액은 15자리 이내로 입력해 주세요.")
    return amount


@dataclass
class SheetSchema:
    name: str
    qr_columns: dict
    name_column: int
    pledge_columns: dict
    date_columns: dict


@dataclass(frozen=True)
class Donation:
    sheet: str
    qr: str
    name: str
    category: str
    year: int
    day: date
    row: int
    coordinate: str
    pledged: object
    existing: object


class ExcelManager:
    def __init__(self, filepath):
        self.filepath = Path(filepath).resolve()
        if self.filepath.suffix.lower() != ".xlsx":
            raise ValueError(".xlsx 파일을 선택해 주세요. .xls는 Excel에서 .xlsx로 변환해야 합니다.")
        self._raw = self.filepath.read_bytes()
        self._fingerprint = sha256(self._raw).digest()
        self._rows, self._indexes, self.changes, self.schemas = {}, {}, {}, {}
        self._formula_cells = {}
        self._sheet_paths = self._read_sheet_paths()
        workbook = openpyxl.load_workbook(BytesIO(self._raw), read_only=True, data_only=True)
        try:
            for sheet in workbook:
                headers = list(sheet.iter_rows(min_row=1, max_row=2, values_only=True))
                schema = self._schema(sheet.title, headers)
                if schema:
                    self.schemas[sheet.title] = schema
        finally:
            workbook.close()
        if not self.schemas:
            raise ValueError("QR 열, 이름, 'YYYY년 작정금액' 머리글이 있는 시트를 찾지 못했습니다.")
        default = next((n for n in self.schemas if n.strip() == "OFF선교헌금"), next(iter(self.schemas)))
        self.select_sheet(default)

    def _read_sheet_paths(self):
        with ZipFile(BytesIO(self._raw)) as archive:
            workbook = ET.fromstring(archive.read("xl/workbook.xml"))
            rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
            targets = {rel.attrib["Id"]: rel.attrib["Target"] for rel in rels}
            paths = {}
            for sheet in workbook.find(f"{{{MAIN_NS}}}sheets"):
                target = targets[sheet.attrib[f"{{{REL_NS}}}id"]]
                paths[sheet.attrib["name"]] = target.lstrip("/") if target.startswith("/") else posixpath.normpath("xl/" + target)
            return paths

    @staticmethod
    def _schema(name, headers):
        if len(headers) < 2:
            return None
        first, second = headers
        qr_columns, pledges, dates = {}, {}, {}
        name_column = None
        for column, value in enumerate(first, 1):
            label = re.sub(r"\s+", "", str(value))
            if label == "이름":
                name_column = column
            for category, title in enumerate(CATEGORIES, 1):
                if label == title + "QR":
                    qr_columns[category] = column
            match = re.fullmatch(r"(\d{4})년작정금액", label)
            block = {}
            for offset in range(5):
                if column + offset <= len(second):
                    sub = str(second[column + offset - 1]).strip()
                    if sub == "개척":
                        sub = "개척교회"
                    if sub in CATEGORIES:
                        block[CATEGORIES.index(sub) + 1] = column + offset
            if len(block) != 5:
                continue
            if match:
                year = int(match[1])
                if year in pledges:
                    raise ValueError(f"{name}: {year}년 작정금액 머리글이 중복되어 있습니다.")
                pledges[year] = block
            if isinstance(value, (date, datetime)):
                day = value.date() if isinstance(value, datetime) else value
                if day in dates:
                    raise ValueError(f"{name}: {day} 날짜 머리글이 중복되어 있습니다.")
                dates[day] = block
        if len(qr_columns) == 5 and name_column and pledges:
            return SheetSchema(name, qr_columns, name_column, pledges, dates)
        return None

    def select_sheet(self, name):
        if name not in self.schemas:
            raise ValueError("지원하는 데이터 시트를 선택해 주세요.")
        if name not in self._rows:
            workbook = openpyxl.load_workbook(BytesIO(self._raw), read_only=True, data_only=True)
            try:
                self._rows[name] = list(workbook[name].values)
            finally:
                workbook.close()
            index = {}
            for row_number, row in enumerate(self._rows[name], 1):
                for category, column in self.schemas[name].qr_columns.items():
                    value = row[column - 1] if column <= len(row) else None
                    if value is None:
                        continue
                    try:
                        qr = parse_qr(value)
                    except ValueError:
                        continue
                    if qr.category == category:
                        index.setdefault(qr.text, []).append(row_number)
            self._indexes[name] = index
            with ZipFile(BytesIO(self._raw)) as archive:
                xml = archive.read(self._sheet_paths[name])
            self._formula_cells[name] = {
                match[1].decode() for match in re.finditer(
                    rb'<c\b[^>]*\br="([A-Z]+\d+)"[^>]*>\s*<f(?:\s|>)', xml
                )
            }
        self.sheet_name = name

    @property
    def schema(self):
        return self.schemas[self.sheet_name]

    @property
    def years(self):
        return sorted(self.schema.pledge_columns)

    def dates(self, year):
        return sorted(day for day in self.schema.date_columns if day.year == int(year))

    @property
    def dirty(self):
        return bool(self.changes)

    def lookup(self, text, year, day):
        qr, year = parse_qr(text), int(year)
        if isinstance(day, datetime):
            day = day.date()
        if year not in self.schema.pledge_columns:
            raise ValueError(f"{year}년 작정금액 열이 없습니다.")
        if day.year != year or day not in self.schema.date_columns:
            raise ValueError(f"{year}년의 입금일 열을 선택해 주세요.")
        matches = self._indexes[self.sheet_name].get(qr.text, [])
        if not matches:
            raise ValueError(f"{qr.text}: '{self.sheet_name}' 전체 행에서 QR을 찾지 못했습니다.")
        if len(matches) != 1:
            raise ValueError(f"{qr.text}: QR이 여러 행에 중복되어 있습니다 ({', '.join(map(str, matches))}).")
        row_number = matches[0]
        row = self._rows[self.sheet_name][row_number - 1]
        column = self.schema.date_columns[day][qr.category]
        coordinate = f"{get_column_letter(column)}{row_number}"
        return Donation(self.sheet_name, qr.text, str(row[self.schema.name_column - 1] or "이름 없음"),
                        CATEGORIES[qr.category - 1], year, day, row_number, coordinate,
                        row[self.schema.pledge_columns[year][qr.category] - 1],
                        self.changes.get((self.sheet_name, coordinate), row[column - 1]))

    def record(self, donation, amount, overwrite=False):
        if donation.sheet != self.sheet_name:
            raise ValueError("시트 선택이 변경되었습니다. QR을 다시 조회해 주세요.")
        current = self.lookup(donation.qr, donation.year, donation.day)
        if current.coordinate != donation.coordinate or current.name != donation.name:
            raise ValueError("선택이 변경되었습니다. QR을 다시 조회해 주세요.")
        amount = parse_amount(amount)
        if current.existing is not None and not overwrite:
            raise ValueError("이 날짜에 이미 금액이 있습니다. 확인 후 변경해 주세요.")
        if donation.coordinate in self._formula_cells[self.sheet_name]:
            raise ValueError("입금 셀에 수식이 있습니다. Excel에서 확인해 주세요.")
        self.changes[(self.sheet_name, donation.coordinate)] = amount
        return amount

    @staticmethod
    def _cell_pattern(coordinate):
        return re.compile(rb'<c\b[^>]*\br="' + coordinate.encode() + rb'"[^>]*?(?:/>|>.*?</c>)', re.DOTALL)

    @classmethod
    def _patch_sheet(cls, xml, changes):
        for coordinate, amount in changes.items():
            match = cls._cell_pattern(coordinate).search(xml)
            if match:
                tag = match[0]
                if re.search(rb"<f(?:\s|>)", tag):
                    raise ValueError(f"{coordinate}: 수식 셀을 덮어쓸 수 없습니다.")
                opening = tag.split(b">", 1)[0].rstrip(b"/")
                opening = re.sub(rb'\s+t="[^"]*"', b"", opening)
                replacement = opening + b'><v>' + str(amount).encode() + b'</v></c>'
                xml = xml[:match.start()] + replacement + xml[match.end():]
            else:
                row_number = re.search(r"\d+$", coordinate)[0]
                pattern = re.compile(rb'<row\b[^>]*\br="' + row_number.encode() + rb'"[^>]*>.*?</row>', re.DOTALL)
                row_match = pattern.search(xml)
                if not row_match:
                    raise ValueError(f"{coordinate}: 원본 행이 없어 기록할 수 없습니다.")
                row = row_match[0]
                column = column_index_from_string(re.match(r"[A-Z]+", coordinate)[0])
                insertion = len(row) - len(b"</row>")
                for cell in re.finditer(rb'<c\b[^>]*\br="([A-Z]+)\d+"', row):
                    if column_index_from_string(cell[1].decode()) > column:
                        insertion = cell.start()
                        break
                cell = f'<c r="{coordinate}"><v>{amount}</v></c>'.encode()
                row = row[:insertion] + cell + row[insertion:]
                xml = xml[:row_match.start()] + row + xml[row_match.end():]
        ET.fromstring(xml)
        return xml

    def save(self):
        if not self.dirty:
            return None
        if sha256(self.filepath.read_bytes()).digest() != self._fingerprint:
            raise ValueError("엑셀 파일이 외부에서 변경되었습니다. 입력 금액을 확인하고 파일을 다시 열어 주세요. 원본에는 덮어쓰지 않았습니다.")
        grouped = {}
        for (sheet, coordinate), amount in self.changes.items():
            grouped.setdefault(self._sheet_paths[sheet], {})[coordinate] = amount
        with ZipFile(BytesIO(self._raw)) as source:
            patches = {path: self._patch_sheet(source.read(path), edits) for path, edits in grouped.items()}
            xml = source.read("xl/workbook.xml")
            match = re.search(rb"<calcPr\b[^>]*(?:/>|>.*?</calcPr>)", xml, re.DOTALL)
            if match:
                calc = match[0].split(b">", 1)[0].rstrip(b"/")
                calc = re.sub(rb'\s+(?:fullCalcOnLoad|forceFullCalc|calcMode)="[^"]*"', b"", calc)
                calc += b' calcMode="auto" fullCalcOnLoad="1" forceFullCalc="1"/>'
                xml = xml[:match.start()] + calc + xml[match.end():]
            else:
                xml = xml.replace(b"</workbook>", b'<calcPr calcMode="auto" fullCalcOnLoad="1" forceFullCalc="1"/></workbook>')
            ET.fromstring(xml)
            patches["xl/workbook.xml"] = xml
            backup_dir = self.filepath.parent / ".backups"
            backup_dir.mkdir(exist_ok=True)
            backup = backup_dir / f"{self.filepath.stem}_{datetime.now():%Y%m%d_%H%M%S_%f}.xlsx"
            shutil.copy2(self.filepath, backup)
            fd, temporary = tempfile.mkstemp(prefix=".qr-save-", suffix=".xlsx", dir=self.filepath.parent)
            try:
                with os.fdopen(fd, "wb") as output:
                    with ZipFile(output, "w") as target:
                        target.comment = source.comment
                        for member in source.infolist():
                            data = patches[member.filename] if member.filename in patches else source.read(member.filename)
                            target.writestr(member, data)
                    output.flush()
                    os.fsync(output.fileno())
                with ZipFile(temporary) as check:
                    if check.testzip() is not None:
                        raise ValueError("저장 파일 검사에 실패했습니다. 원본은 보존되었습니다.")
                if sha256(self.filepath.read_bytes()).digest() != self._fingerprint:
                    raise ValueError("저장 중 원본이 변경되었습니다. 원본을 보존했습니다.")
                os.replace(temporary, self.filepath)
            finally:
                if os.path.exists(temporary):
                    os.unlink(temporary)
        for (sheet, coordinate), amount in self.changes.items():
            column, number = coordinate_from_string(coordinate)
            row = list(self._rows[sheet][number - 1])
            row[column_index_from_string(column) - 1] = amount
            self._rows[sheet][number - 1] = tuple(row)
        self._raw = self.filepath.read_bytes()
        self._fingerprint = sha256(self._raw).digest()
        self.changes.clear()
        return backup

    def close(self):
        self._rows.clear()
        self._indexes.clear()
        self._formula_cells.clear()
